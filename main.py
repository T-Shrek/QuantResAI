import json
import os
import time
from pathlib import Path
from docx import Document
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field
from pypdf import PdfReader

load_dotenv()

API_KEY=os.getenv("GROQ_API_KEY")
if not API_KEY:
    raise ValueError("GROQ_API_KEY environment variable not set.")

client=Groq(api_key=API_KEY)
MODEL_NAME="openai/gpt-oss-20b"

SKILL_ALIASES={
    # Programming languages
    "js": "javascript",
    "javascript": "javascript",
    "ts": "typescript",
    "typescript": "typescript",
    "py": "python",
    "python": "python",
    "cpp": "c++",
    "c plus plus": "c++",
    # Backend
    "node": "node.js",
    "nodejs": "node.js",
    "node.js": "node.js",
    "expressjs": "express",
    "express.js": "express",
    # Frontend
    "reactjs": "react",
    "react.js": "react",
    "react": "react",
    # Databases
    "postgres": "postgresql",
    "postgresql": "postgresql",
    "mysql": "mysql",
    "mongo": "mongodb",
    "mongodb": "mongodb",
    # Cloud
    "aws": "aws",
    "amazon web services": "aws",
    "gcp": "gcp",
    "google cloud": "gcp",
    "azure": "azure",
    # DevOps
    "k8s": "kubernetes",
    "kubernetes": "kubernetes",
    "docker": "docker",
    # Other
    "rest api": "rest",
    "rest apis": "rest",
    "restful api": "rest",
    "restful apis": "rest"
}
WEIGHTS={
    "required_skills": 40,
    "preferred_skills": 15,
    "experience": 20,
    "education": 10,
    "projects": 10,
    "certifications": 5
}
class JobD(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float|None
    education_requirements: list[str]
    responsibilities: list[str]

class Experience(BaseModel):
    company: str|None=None
    role: str|None=None
    duration: str|None=None
    description: str|None=None
    skills_used: list[str]=Field(default_factory=list)

class Resume(BaseModel):
    name: str|None=None
    email: str|None=None
    phone: str|None=None
    total_experience_years: float|None=None
    skills: list[str]=Field(default_factory=list)
    experience: list[Experience]=Field(default_factory=list)
    education: list[str]=Field(default_factory=list)
    projects: list[str]=Field(default_factory=list)
    certifications: list[str]=Field(default_factory=list)

class MatchResult(BaseModel):
    score: float
    details: dict

def read_pdf(file_path: Path)->str:
    reader=PdfReader(file_path)
    text=[]
    for page in reader.pages:
        page_text=page.extract_text()
        if page_text:
            text.append(page_text)
    return "\n".join(text)

def read_docx(file_path: Path) -> str:
    doc=Document(file_path)
    text=[
        p.text.strip()
        for p in doc.paragraphs
        if p.text.strip()
    ]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text.append(cell.text.strip())
    return "\n".join(text)

def extract_resume_text(file_path: Path) -> str|None:
    suffix=file_path.suffix.lower()
    if suffix == ".pdf":
        return read_pdf(file_path)
    elif suffix == ".docx":
        return read_docx(file_path)
    return None

def normalize_skill(skill: str) -> str:
    skill=skill.strip().lower()
    return SKILL_ALIASES.get(skill, skill)

def normalize_skills(skills: list[str]) -> list[str]:
    normalized=[]
    for skill in skills:
        skill=normalize_skill(skill)
        if skill not in normalized:
            normalized.append(skill)
    return normalized

def parse_job_description(jd_text: str) -> JobD:
    schema=JobD.model_json_schema()
    system_prompt=(
        "You are an expert HR assistant. Analyze job descriptions and extract structured details. "
        f"Return ONLY valid JSON strictly adhering to this schema:\n{schema}"
    )
    response=client.chat.completions.create(
        model=MODEL_NAME,
        messages=[
            {
                "role": "system",
                "content": system_prompt
            },
            {
                "role": "user",
                "content": jd_text
            },
        ],
        response_format={"type": "json_object"},
    )
    data=json.loads(
        response.choices[0].message.content
    )
    job=JobD(**data)
    job.required_skills=normalize_skills(
        job.required_skills
    )
    job.preferred_skills=normalize_skills(
        job.preferred_skills
    )
    return job

def parse_resume(resume_text: str)->Resume:
    schema=Resume.model_json_schema()
    system_prompt=(
        "You are an expert ATS parser. Extract information based on semantic meaning, "
        "not just exact heading names. Include internships in experience. "
        f"Return ONLY valid JSON matching this schema:\n{schema}"
    )
    response=client.chat.completions.create(
        model=MODEL_NAME,
        messages=[
            {
                "role": "system",
                "content": system_prompt
            },
            {
                "role": "user",
                "content": resume_text
            },
        ],
        response_format={"type": "json_object"},
    )
    data=json.loads(
        response.choices[0].message.content
    )
    resume=Resume(**data)
    resume.skills=normalize_skills(
        resume.skills
    )
    for experience in resume.experience:
        experience.skills_used=normalize_skills(
            experience.skills_used
        )
    return resume

def calculate_skill_match(
    required_skills,
    resume_skills
):
    required=set(
        normalize_skills(required_skills)
    )
    resume=set(
        normalize_skills(resume_skills)
    )
    if not required:
        return 100.0, [], []
    matching=required.intersection(resume)
    missing=required - resume
    score=(
        len(matching)/len(required)
    ) * 100
    return (
        score,
        list(matching),
        list(missing)
    )

def calculate_preferred_skill_match(
    preferred_skills,
    resume_skills
):
    preferred=set(
        normalize_skills(preferred_skills)
    )
    resume=set(
        normalize_skills(resume_skills)
    )
    if not preferred:
        return 100.0
    matching=preferred.intersection(resume)
    return (
        len(matching)/len(preferred)
    ) * 100

def calculate_experience_score(
    candidate_years,
    required_years
):
    if required_years is None:
        return 100.0
    if candidate_years is None:
        return 0.0
    if candidate_years >= required_years:
        return 100.0
    return (
        candidate_years / required_years
    ) * 100
def calculate_education_score(
    resume,
    job
):
    if not job.education_requirements:
        return 100.0
    if not resume.education:
        return 0.0
    education_text=" ".join(
        resume.education
    ).lower()
    matches=0
    for requirement in job.education_requirements:
        if requirement.lower() in education_text:
            matches += 1
    return (
        matches / len(job.education_requirements)
    ) * 100

def calculate_project_score(resume):
    if not resume.projects:
        return 0.0
    if len(resume.projects) >= 2:
        return 100.0
    return 50.0

def calculate_certification_score(resume):
    if resume.certifications:
        return 100.0
    return 0.0

def calculate_weighted_score(job, resume):
    required_score, matching_skills, missing_skills=calculate_skill_match(
        job.required_skills,
        resume.skills
    )
    preferred_score=calculate_preferred_skill_match(
        job.preferred_skills,
        resume.skills
    )
    experience_score=calculate_experience_score(
        resume.total_experience_years,
        job.minimum_experience
    )
    education_score=calculate_education_score(
        resume,
        job
    )
    project_score=calculate_project_score(
        resume
    )
    certification_score=calculate_certification_score(
        resume
    )
    final_score=(
        required_score * WEIGHTS["required_skills"] / 100
        + preferred_score * WEIGHTS["preferred_skills"] / 100
        + experience_score * WEIGHTS["experience"] / 100
        + education_score * WEIGHTS["education"] / 100
        + project_score * WEIGHTS["projects"] / 100
        + certification_score * WEIGHTS["certifications"] / 100
    )
    return {
        "score": round(final_score, 2),
        "breakdown": {
            "required_skills": round(required_score, 2),
            "preferred_skills": round(preferred_score, 2),
            "experience": round(experience_score, 2),
            "education": round(education_score, 2),
            "projects": round(project_score, 2),
            "certifications": round(certification_score, 2)
        },
        "matching_skills": matching_skills,
        "missing_skills": missing_skills
    }

def evaluate_candidate(
    job: JobD,
    resume: Resume
) -> MatchResult:
    score_result=calculate_weighted_score(
        job,
        resume
    )
    prompt=f"""
You are an HR recruiter evaluating a candidate.
JOB REQUIREMENTS:
{job.model_dump_json(indent=2)}
CANDIDATE RESUME:
{resume.model_dump_json(indent=2)}
DETERMINISTIC SCORE:
{json.dumps(score_result, indent=2)}
The score has already been calculated by the
recruitment scoring system.
DO NOT change the score.
Use the score and breakdown to explain the candidate.
Return ONLY a JSON object with this shape:

{{
    "score": {score_result["score"]},
    "details": {{
        "candidate_name": "{resume.name}",
        "matching_skills": {json.dumps(score_result["matching_skills"])},
        "missing_skills": {json.dumps(score_result["missing_skills"])},
        "score_breakdown": {json.dumps(score_result["breakdown"])},
        "experience_requirement_met": true,
        "verdict": "Short summary of evaluation"
    }}
}}
"""
    response=client.chat.completions.create(
        model=MODEL_NAME,
        messages=[
            {
                "role": "user",
                "content": prompt
            }
        ],
        response_format={"type": "json_object"},
    )
    data=json.loads(
        response.choices[0].message.content
    )
    return MatchResult(**data)

def main():
    jd_path=Path("job_description.txt")
    resume_folder=Path("resumes")
    if not jd_path.exists():
        raise FileNotFoundError(
            f"Job description file '{jd_path}' not found."
        )
    if not resume_folder.exists():
        raise FileNotFoundError(
            f"Directory '{resume_folder}' not found."
        )
    
    print("Parsing Job Description...")
    job_description_text=jd_path.read_text(
        encoding="utf-8"
    )
    parsed_job=parse_job_description(
        job_description_text
    )
    results=[]
    for file_path in resume_folder.iterdir():
        if file_path.suffix.lower() not in [
            ".pdf",
            ".docx"
        ]:
            continue
        print(
            f"Parsing resume: {file_path.name}"
        )
        try:
            resume_text=extract_resume_text(
                file_path
            )
            if not resume_text:
                continue
            parsed_resume=parse_resume(
                resume_text
            )
            time.sleep(2)
            evaluation=evaluate_candidate(
                parsed_job,
                parsed_resume
            )
            time.sleep(2)
            results.append({
                "name": parsed_resume.name or file_path.stem,
                "score": evaluation.score,
                "details": evaluation.details,
            })
            print(
                f"Finished {file_path.name} "
                f"- Score: {evaluation.score}"
            )
            print(
                "Score Breakdown:",
                evaluation.details.get(
                    "score_breakdown",
                    {}
                )
            )
            print(
                "Matching Skills:",
                evaluation.details.get(
                    "matching_skills",
                    []
                )
            )
            print(
                "Missing Skills:",
                evaluation.details.get(
                    "missing_skills",
                    []
                )
            )
        except Exception as e:
            print(
                f"Error processing "
                f"{file_path.name}: {e}"
            )
    results.sort(
        key=lambda x: x["score"],
        reverse=True
    )
    print("\n--- TOP CANDIDATES ---")
    for candidate in results[:2]:
        print(
            f"{candidate['name']} - "
            f"Score: {candidate['score']}"
        )
    print("\n--- BOTTOM CANDIDATES ---")
    for candidate in results[-2:]:
        print(
            f"{candidate['name']} - "
            f"Score: {candidate['score']}"
        )
        
if __name__ == "__main__":
    main()